"""Document parsers for PDF, DOCX, TXT, MD."""

from __future__ import annotations

from pathlib import Path

from ..types import Document


class UnsupportedFormatError(Exception):
    """Raised when file format is not supported."""


def parse_file(file_path: str, doc_type: str | None = None) -> Document:
    """Parse a file into a Document node.

    Args:
        file_path: Path to the file.
        doc_type: Override format detection ('pdf', 'docx', 'txt', 'md').

    Raises:
        FileNotFoundError: If file does not exist.
        UnsupportedFormatError: If format is not recognized.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if doc_type is None:
        doc_type = path.suffix.lstrip(".").lower()

    if doc_type == "pdf":
        content = _parse_pdf(str(path))
    elif doc_type == "docx":
        content = _parse_docx(str(path))
    elif doc_type == "txt":
        content = _parse_txt(str(path))
    elif doc_type in ("md", "markdown"):
        content = _parse_md(str(path))
    else:
        raise UnsupportedFormatError(f"Unsupported format: {doc_type}")

    return Document(
        content=content,
        source_path=str(path.resolve()),
        doc_type=doc_type,
        metadata={"filename": path.name, "size_bytes": path.stat().st_size},
    )


def _parse_pdf(path: str) -> str:
    import pymupdf

    doc = pymupdf.open(path)
    parts: list[str] = []
    for page in doc:
        parts.append(page.get_text())
    doc.close()
    return "\n".join(parts)


def _parse_docx(path: str) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    return "\n".join(para.text for para in doc.paragraphs if para.text.strip())


def _parse_txt(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


def _parse_md(path: str) -> str:
    import markdown
    from html.parser import HTMLParser

    class _TextExtractor(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self._parts: list[str] = []

        def handle_data(self, data: str) -> None:
            self._parts.append(data)

        def get_text(self) -> str:
            return "".join(self._parts)

    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    html = markdown.markdown(raw)
    extractor = _TextExtractor()
    extractor.feed(html)
    return extractor.get_text()